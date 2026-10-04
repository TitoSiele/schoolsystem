from __future__ import annotations

import io
import math
import re
from pathlib import Path
from typing import Any

FIELD_ALIASES = {
    "admission_number": {"admno", "admissionno", "admissionnumber", "studentid", "studentnumber", "registrationnumber", "regno"},
    "student_name": {"studentname", "name", "learnername", "pupilname"},
    "first_name": {"firstname", "givenname", "forenames"},
    "last_name": {"lastname", "surname", "familyname"},
    "gender": {"gender", "sex"},
    "class_name": {"class", "classname", "grade", "form", "stream", "level"},
    "parent_name": {"parentname", "guardianname", "guardian", "nextofkin"},
    "parent_phone": {"parentphone", "guardianphone", "phone", "phonenumber", "telephone", "mobile"},
    "fee_balance": {"feebalance", "balance", "amountdue", "outstanding", "outstandingbalance", "feesdue", "schoolfees"},
}


def clean_text(value: Any) -> str | None:
    if value is None:
        return None
    # pandas/openpyxl turn blank spreadsheet cells into NaN, and str(NaN) is the
    # literal string "nan". Without this, an empty Admission No cell becomes the
    # text "nan" and is imported as a real value.
    if isinstance(value, float) and math.isnan(value):
        return None
    text = re.sub(r"\s+", " ", str(value).replace("\n", " ").replace("\r", " ")).strip()
    if not text or text.lower() in {"nan", "nat", "none", "null"}:
        return None
    return text


def normalize_header(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", "", (clean_text(value) or "").lower())


def map_columns(columns: list[Any]) -> dict[str, Any]:
    mapped: dict[str, Any] = {}
    aliases = {alias: field for field, values in FIELD_ALIASES.items() for alias in values}
    for column in columns:
        normalized = normalize_header(column)
        field = aliases.get(normalized)
        if field and field not in mapped:
            mapped[field] = column
    return mapped


def split_student_name(value: Any) -> tuple[str, str]:
    parts = (clean_text(value) or "").split(" ", 1)
    return parts[0] if parts else "", parts[1] if len(parts) > 1 else ""


def normalize_records(records: list[dict[str, Any]], source: str) -> list[dict[str, Any]]:
    normalized_records = []
    for raw in records:
        column_map = map_columns(list(raw.keys()))
        item: dict[str, Any] = {"source": source}
        for field, column in column_map.items():
            item[field] = clean_text(raw.get(column))
        if not item.get("first_name") and item.get("student_name"):
            item["first_name"], item["last_name"] = split_student_name(item["student_name"])
        if item.get("fee_balance"):
            number = re.sub(r"[^0-9.\-]", "", item["fee_balance"] or "")
            try:
                item["fee_balance"] = float(number) if number else 0.0
            except ValueError:
                item["fee_balance"] = None
        missing = [field for field in ("admission_number", "first_name", "last_name") if not item.get(field)]
        item["error_flags"] = missing
        item["confidence_score"] = round(max(0.0, 1.0 - (len(missing) * 0.25)), 3)
        normalized_records.append(item)
    return normalized_records


def extract_spreadsheet(content: bytes, suffix: str) -> tuple[list[dict[str, Any]], str]:
    import pandas as pd

    engine = "openpyxl" if suffix == ".xlsx" else None
    workbook = pd.read_excel(io.BytesIO(content), engine=engine) if suffix == ".xlsx" else pd.read_csv(io.BytesIO(content))
    workbook = workbook.where(workbook.notna(), None)
    return workbook.to_dict(orient="records"), "spreadsheet"


def extract_pdf(content: bytes) -> tuple[list[dict[str, Any]], str]:
    import pdfplumber

    rows: list[dict[str, Any]] = []
    text_blocks: list[str] = []
    with pdfplumber.open(io.BytesIO(content)) as pdf:
        for page in pdf.pages:
            tables = page.extract_tables() or []
            for table in tables:
                if table:
                    headers = table[0]
                    rows.extend(dict(zip(headers, row)) for row in table[1:] if row and any(row))
            if not tables:
                text = page.extract_text() or ""
                text_blocks.extend(line for line in text.splitlines() if line.strip())
    if not rows:
        rows = parse_text_rows(text_blocks)
    return rows, "pdf"


def extract_docx(content: bytes) -> tuple[list[dict[str, Any]], str]:
    """Word documents (.docx). Word tables are parsed first because a school
    roster is almost always laid out as a table, not flowing paragraphs."""
    import docx

    document = docx.Document(io.BytesIO(content))

    rows: list[list[str]] = []
    for table in document.tables:
        for row in table.rows:
            cells = [(cell.text or "").strip() for cell in row.cells]
            if any(cells):
                rows.append(cells)

    if not rows:
        # Fall back to paragraph text for a plain list-style document.
        lines = [p.text for p in document.paragraphs]
        rows = parse_text_rows([ln for ln in lines if ln.strip()])

    if not rows:
        return [], "word"

    headers = rows[0]
    records = [dict(zip(headers, row)) for row in rows[1:]]
    return records, "word"


def extract_image(content: bytes) -> tuple[list[dict[str, Any]], str]:
    """OCR for scanned/image documents.

    pytesseract is only a Python wrapper - it still needs the Tesseract binary
    installed on the machine. When that binary is missing we raise a clear,
    actionable error rather than an opaque one.
    """
    from PIL import Image
    import pytesseract

    try:
        pytesseract.get_tesseract_version()
    except Exception as exc:  # binary missing from PATH
        raise ValueError(
            "Image/OCR support needs the Tesseract engine installed on this machine. "
            "Install Tesseract OCR and add it to PATH, or upload a CSV/XLSX/PDF/DOCX instead."
        ) from exc

    text = pytesseract.image_to_string(Image.open(io.BytesIO(content)))
    return parse_text_rows(text.splitlines()), "image_ocr"


def parse_text_rows(lines: list[str]) -> list[dict[str, Any]]:
    useful = [re.sub(r"\s+", " ", line).strip() for line in lines if line.strip()]
    if not useful:
        return []
    header_parts = re.split(r"\s{2,}|\t|,|;", useful[0])
    if len(header_parts) >= 2 and map_columns(header_parts):
        headers = [part.strip() for part in header_parts]
        data_lines = useful[1:]
    else:
        headers = ["admission_number", "student_name", "class_name", "fee_balance"]
        data_lines = useful
    rows = []
    for line in data_lines:
        parts = [part.strip() for part in re.split(r"\s{2,}|\t|,|;", line)]
        if len(parts) < 2:
            continue
        rows.append(dict(zip(headers, parts)))
    return rows


def extract_document(content: bytes, filename: str, content_type: str | None) -> tuple[list[dict[str, Any]], str]:
    suffix = Path(filename).suffix.lower()
    if suffix == ".csv" or content_type == "text/csv":
        records, source = extract_spreadsheet(content, ".csv")
    elif suffix in {".xlsx", ".xls"}:
        records, source = extract_spreadsheet(content, suffix)
    elif suffix == ".docx" or content_type in {
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    }:
        records, source = extract_docx(content)
    elif suffix == ".pdf" or content_type == "application/pdf":
        records, source = extract_pdf(content)
    elif suffix in {".jpg", ".jpeg", ".png", ".tif", ".tiff", ".bmp"} or (content_type or "").startswith("image/"):
        records, source = extract_image(content)
    else:
        raise ValueError("Unsupported document type. Use CSV, XLSX, DOCX, PDF, JPG, PNG, TIFF, or BMP.")
    return normalize_records(records, source), source
